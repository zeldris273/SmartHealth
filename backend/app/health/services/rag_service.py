from __future__ import annotations

import io
import re
from dataclasses import dataclass

from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader
from sqlalchemy.orm import Session

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from openai import OpenAI
except ImportError:
    OpenAI = None

from app.health.core.config import settings
from app.health.models import KnowledgeChunk, KnowledgeDocument


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 180


@dataclass
class RetrievedChunk:
    document_id: int
    filename: str
    chunk_index: int
    content: str
    score: float


def get_openai_client() -> OpenAI:
    api_key = getattr(settings, "OPENAI_API_KEY", None) or getattr(settings, "OPEN_API_KEY", None)
    if not api_key or api_key in {"your-openai-api-key-here", ""}:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="OPENAI_API_KEY chưa được cấu hình trong file .env.",
        )
    if OpenAI is None:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Chưa cài thư viện openai. Hãy chạy: pip install -r requirements.txt",
        )
    return OpenAI(api_key=api_key)


def get_file_extension(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot >= 0 else ""


def normalize_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_txt(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx(content: bytes) -> str:
    if DocxDocument is None:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Chưa cài python-docx để đọc DOCX.",
        )
    try:
        document = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Định dạng file DOCX không hợp lệ hoặc bị lỗi: {exc}",
        )
    
    text_parts = []
    # Extract text from paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)
            
    # Extract text from tables
    for table in document.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                text_parts.append(" | ".join(row_text))
                
    return "\n".join(text_parts)


HEALTH_AND_HOSPITAL_KEYWORDS = {
    "sức khỏe", "suc khoe", "y tế", "y te", "y khoa", "y học", "y hoc",
    "bệnh viện", "benh vien", "phòng khám", "phong kham", "trạm y tế", "tram y te",
    "bệnh xá", "benh xa", "nhà thuốc", "nha thuoc", "quầy thuốc", "quay thuoc",
    "bác sĩ", "bac si", "y tá", "y ta", "điều dưỡng", "dieu duong", "dược sĩ", "duoc si",
    "nha sĩ", "nha si", "hộ sinh", "ho sinh", "thầy thuốc", "thay thuoc",
    "bệnh", "benh", "dịch bệnh", "dich benh", "triệu chứng", "trieu chung",
    "đau", "dau", "sốt", "sot", "ho", "sổ mũi", "so mui", "chóng mặt", "chong mat",
    "nhức đầu", "nhuc dau", "mệt mỏi", "met moi", "buồn nôn", "buon non",
    "dị ứng", "di ung", "nhiễm trùng", "nhiem trung", "viêm", "viem",
    "thuốc", "thuoc", "vắc xin", "vac xin", "vaccine", "kê đơn", "ke don",
    "khám", "kham", "điều trị", "dieu tri", "chữa bệnh", "chua benh",
    "phẫu thuật", "phau thuat", "mổ", "mo", "cấp cứu", "cap cuu",
    "xét nghiệm", "xet nghiem", "chụp x-quang", "chup x-quang", "siêu âm", "sieu am",
    "bmi", "cân nặng", "can nang", "chiều cao", "chieu cao", "calo", "calorie", "calories",
    "dinh dưỡng", "dinh duong", "chế độ ăn", "che do an", "kiêng", "kieng",
    "tập luyện", "tap luyen", "thể dục", "the duc", "thể thao", "the thao",
    "giảm cân", "giam can", "tăng cân", "tang can", "tăng cơ", "tang co",
    "protein", "carb", "chất béo", "chat beo", "tdee", "bmr",
    "tim", "huyết áp", "huyet ap", "đường huyết", "duong huyet", "tiểu đường", "tieu duong",
    "cholesterol", "vitamin", "khoáng chất", "khoang chat",
}


def is_document_health_related(text: str) -> bool:
    normalized = text.lower()
    matched = [kw for kw in HEALTH_AND_HOSPITAL_KEYWORDS if kw in normalized]
    if len(matched) >= 2:
        return True
    if len(matched) == 1 and len(normalized) < 300:
        return True

    api_key = getattr(settings, "OPENAI_API_KEY", None) or getattr(settings, "OPEN_API_KEY", None)
    if api_key and api_key not in {"your-openai-api-key-here", ""}:
        try:
            client = get_openai_client()
            sample_text = text[:2000]
            response = client.chat.completions.create(
                model=getattr(settings, "OPENAI_MODEL", "gpt-4.1-mini"),
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "Bạn là trợ lý phân loại tài liệu cho hệ thống SmartHealth. "
                            "Hãy xác định xem tài liệu được cung cấp dưới đây có liên quan đến chủ đề sức khỏe, "
                            "y tế, bệnh tật, bệnh viện, phòng khám, bác sĩ, dinh dưỡng, thể dục, chỉ số cơ thể, hoặc "
                            "các vấn đề liên quan đến y tế/sức khỏe hay không.\n"
                            "Chỉ trả lời chính xác 'YES' hoặc 'NO'. Không giải thích gì thêm."
                        )
                    },
                    {
                        "role": "user",
                        "content": f"Nội dung tài liệu:\n{sample_text}"
                    }
                ],
                temperature=0.0,
                max_tokens=5,
            )
            result = response.choices[0].message.content.strip().upper()
            if "YES" in result:
                return True
        except Exception:
            return len(matched) > 0

    return False


def extract_document_text(filename: str, content: bytes) -> str:
    extension = get_file_extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Chỉ hỗ trợ upload PDF, DOCX và TXT.",
        )

    if extension == ".pdf":
        text = extract_pdf(content)
    elif extension == ".docx":
        text = extract_docx(content)
    else:
        text = extract_txt(content)

    text = normalize_text(text)
    if not text:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Không trích xuất được nội dung từ tài liệu.",
        )
    return text


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        if len(paragraph) > chunk_size:
            if current:
                chunks.append(current.strip())
                current = ""
            start = 0
            while start < len(paragraph):
                chunks.append(paragraph[start : start + chunk_size].strip())
                start += chunk_size - overlap
            continue

        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= chunk_size:
            current = candidate
        else:
            chunks.append(current.strip())
            current = paragraph

    if current:
        chunks.append(current.strip())

    return [chunk for chunk in chunks if chunk]


def create_embeddings(texts: list[str]) -> list[list[float]]:
    if not texts:
        return []

    client = get_openai_client()
    response = client.embeddings.create(
        model=getattr(settings, "OPENAI_EMBEDDING_MODEL", "text-embedding-3-small"),
        input=texts,
    )
    return [item.embedding for item in response.data]


async def ingest_upload(db: Session, user_id: int, file: UploadFile) -> KnowledgeDocument:
    filename = file.filename or "document"
    content = await file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="File quá lớn. Giới hạn hiện tại là 10MB.",
        )

    text = extract_document_text(filename, content)
    
    if not is_document_health_related(text):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Tài liệu tải lên không liên quan đến chủ đề y tế, sức khỏe hoặc bệnh viện. Hệ thống chỉ hỗ trợ các tài liệu thuộc lĩnh vực này.",
        )

    chunks = chunk_text(text)
    embeddings = create_embeddings(chunks)

    document = KnowledgeDocument(
        user_id=user_id,
        filename=filename,
        content_type=file.content_type,
        status="processed",
        chunk_count=len(chunks),
    )
    db.add(document)
    db.flush()

    for index, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
        db.add(
            KnowledgeChunk(
                document_id=document.id,
                user_id=user_id,
                chunk_index=index,
                content=chunk,
                embedding=embedding,
            )
        )

    db.commit()
    db.refresh(document)
    return document


def search_relevant_chunks(
    db: Session,
    user_id: int,
    query: str,
    limit: int | None = None,
) -> list[RetrievedChunk]:
    has_chunks = (
        db.query(KnowledgeChunk.id)
        .filter(
            KnowledgeChunk.user_id == user_id,
            KnowledgeChunk.is_deleted == False,  # noqa: E712
        )
        .first()
    )
    if not has_chunks:
        return []

    embedding = create_embeddings([query])[0]
    top_k = limit or getattr(settings, "RAG_TOP_K", 5)
    distance = KnowledgeChunk.embedding.cosine_distance(embedding)

    rows = (
        db.query(
            KnowledgeChunk,
            KnowledgeDocument.filename,
            distance.label("distance"),
        )
        .join(KnowledgeDocument, KnowledgeDocument.id == KnowledgeChunk.document_id)
        .filter(
            KnowledgeChunk.user_id == user_id,
            KnowledgeChunk.is_deleted == False,  # noqa: E712
            KnowledgeDocument.is_deleted == False,  # noqa: E712
        )
        .order_by(distance)
        .limit(top_k)
        .all()
    )

    threshold = getattr(settings, "RAG_SIMILARITY_THRESHOLD", 0.35)
    ranked = [
        RetrievedChunk(
            document_id=chunk.document_id,
            filename=filename,
            chunk_index=chunk.chunk_index,
            content=chunk.content,
            score=round(1 - float(distance), 4),
        )
        for chunk, filename, distance in rows
    ]
    if not ranked:
        return []

    filtered = [chunk for chunk in ranked if chunk.score >= threshold]
    return filtered


def format_retrieved_context(chunks: list[RetrievedChunk]) -> str:
    if not chunks:
        return "Không tìm thấy chunk tài liệu liên quan."

    return "\n\n".join(
        (
            f"[Tài liệu: {chunk.filename} | chunk {chunk.chunk_index + 1} | "
            f"độ liên quan {chunk.score}]\n{chunk.content}"
        )
        for chunk in chunks
    )
