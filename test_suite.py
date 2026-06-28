from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
import time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

EMAIL = "sekaikamiki2309@gmail.com"
PASSWORD = "Roti2324@232400"

driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()))
wait = WebDriverWait(driver, 20)

BASE_URL = "http://localhost:5173"

# =========================
# EXCEL REPORT
# =========================

wb = Workbook()
ws = wb.active
ws.title = "Test Report"

ws.append(["Test Case", "Status", "Detail"])

green = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
red = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")

for cell in ws[1]:
    cell.font = Font(bold=True)

pass_count = 0
fail_count = 0

def report(tc, status, detail):
    global pass_count, fail_count

    ws.append([tc, status, detail])

    row = ws.max_row

    if status == "PASS":
        ws[f"B{row}"].fill = green
        pass_count += 1
    else:
        ws[f"B{row}"].fill = red
        fail_count += 1

# =========================
# UTILS
# =========================
def log(msg):
    print("\n====================")
    print(msg)
    print("====================\n")

def click(xpath):
    el = wait.until(EC.element_to_be_clickable((By.XPATH, xpath)))
    driver.execute_script("arguments[0].click();", el)
    return el

def safe_get(url):
    driver.get(url)
    time.sleep(2)

# =========================
# 1. OPEN HOME
# =========================
log("TEST 1: OPEN HOME PAGE")

safe_get(BASE_URL)

if "Smart Health" in driver.title:
    print("✅ Home loaded")
    report("Home Page","PASS","Home page loaded")
else:
    print("❌ Home failed")
    report("Home Page","FAIL","Cannot load home page")

# =========================
# 2. OPEN LOGIN MODAL
# =========================
log("TEST 2: OPEN LOGIN MODAL")

login_btn = wait.until(
    EC.element_to_be_clickable((By.CSS_SELECTOR, "button"))
)
driver.execute_script("arguments[0].click();", login_btn)

time.sleep(2)

inputs = wait.until(
    EC.presence_of_all_elements_located((By.CSS_SELECTOR, "input"))
)

if len(inputs) >= 2:
    print("✅ Login modal opened")
    report("Login Modal","PASS","Modal displayed")
else:
    print("❌ Login modal failed")
    report("Login Modal","FAIL","Modal not displayed")

# =========================
# 3. LOGIN
# =========================
log("TEST 3: LOGIN FUNCTION")

inputs[0].clear()
inputs[0].send_keys(EMAIL)

inputs[1].clear()
inputs[1].send_keys(PASSWORD)

click("//button[contains(., 'Đăng nhập') or @type='submit']")

time.sleep(4)

if "/dashboard" in driver.current_url or "/profile" in driver.current_url:
    print("✅ Login success")
    report("Login", "PASS", "Login successful")
else:
    print("❌ Login failed")
    report("Login", "FAIL", "Cannot login")
# =========================
# 4. TEST NAVIGATION
# =========================
log("TEST 4: NAVIGATION TEST")

pages = [
    ("/dashboard", "Dashboard"),
    ("/profile", "Profile"),
    ("/chat", "Chat"),
]

for path, name in pages:
    safe_get(BASE_URL + path)

    if name.lower() in driver.page_source.lower():
        print(f"✅ {name} page OK")
        report(name,"PASS","Opened successfully")
    else:
        print(f"⚠️ {name} may require auth or failed")
        report(name,"FAIL","Cannot open page")

# =========================
# 5. TEST HEADER UI
# =========================
log("TEST 5: HEADER ELEMENTS")

safe_get(BASE_URL)

header_checks = [
    "//header",
    "//button"
]

for x in header_checks:
    try:
        driver.find_element(By.XPATH, x)
        print(f"✅ Found: {x}")
        report("Header","PASS",x)
    except:
        print(f"❌ Missing: {x}")
        report("Header","FAIL",x)
# =========================
# 6. TEST CSKH WIDGET EXISTS
# =========================
log("TEST 6: CSKH WIDGET")

if "Mở chat CSKH" in driver.page_source:
    print("✅ CSKH widget exists")
    report("CSKH Widget","PASS","Widget found")
else:
    print("⚠️ CSKH widget not found")
    report("CSKH Widget","FAIL","Widget missing")

# =========================
# 7. PERFORMANCE CHECK (basic)
# =========================
# =========================
# 7. PERFORMANCE CHECK (basic)
# =========================
log("TEST 7: BASIC LOAD CHECK")

start = time.time()

driver.get(BASE_URL)

end = time.time()

load_time = round(end - start, 2)

print(f"Page load time: {load_time}s")

if load_time < 5:
    print("✅ Performance PASS")
    report("Performance", "PASS", f"Load time: {load_time}s")
else:
    print("❌ Performance FAIL")
    report("Performance", "FAIL", f"Load time: {load_time}s")

# =========================
# FINISH
# =========================
log("ALL TEST DONE")

time.sleep(3)

# =========================
# SAVE EXCEL REPORT
# =========================

ws.append([])
ws.append(["TOTAL PASS", pass_count])
ws.append(["TOTAL FAIL", fail_count])
ws.append(["TOTAL TEST", pass_count + fail_count])

wb.save("Test_Report.xlsx")

print("\n==============================")
print("Excel report saved successfully!")
print("File: Test_Report.xlsx")
print(f"PASS: {pass_count}")
print(f"FAIL: {fail_count}")
print("==============================")

driver.quit()